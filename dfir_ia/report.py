"""Salida JSON, CSV, Markdown y opcional DOCX; escape de datos no confiables."""
from __future__ import annotations
import csv
import html
import json
from pathlib import Path
import re
from datetime import datetime, timezone


def _safe(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(s))

def _csvsafe(value):
    text = _safe(value)
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) else text


def report(output, findings, events, manifest, warnings, ai=None, docx=True):
    dst = Path(output)
    dst.mkdir(parents=True, exist_ok=True)
    hosts = sorted({e.computer for e in events} | {e['evidence'][0]['computer'] for e in findings if e['evidence']})
    data = {
        'schema_version': '1.0', 'generated_utc': datetime.now(timezone.utc).isoformat(),
        'status': 'Informe preliminar sujeto a validación forense',
        'event_count': len(events), 'hosts': hosts,
        'findings': findings, 'source_manifest': manifest, 'warnings': warnings,
        'ai_review': ai,
    }
    # Generar DOCX antes que los otros formatos permite conservar los avisos
    # de dependencias ausentes dentro de hallazgos.json e informes HTML/MD.
    if docx:
        try:
            _make_docx(dst / 'informe.docx', data)
        except ImportError:
            warnings.append('No se generó DOCX: instalar python-docx (pip install -r requirements.txt).')
    (dst / 'hallazgos.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    with (dst / 'timeline.csv').open('w', newline='', encoding='utf-8-sig') as fp:
        w = csv.writer(fp)
        w.writerow(['Timestamp', 'Computer', 'EventID', 'Channel', 'RecordID', 'Level', 'RuleTitle', 'Details', 'Source'])
        for e in sorted(events, key=lambda e: (e.dt is None, e.dt or datetime.max.replace(tzinfo=timezone.utc))):
            w.writerow([_csvsafe(x) for x in (e.timestamp, e.computer, e.event_id, e.channel,
                        e.record_id, e.level, e.title, e.details, e.source)])
    lines = [
        '# Informe de análisis forense automatizado (PRELIMINAR)', '',
        '**Estado:** Pendiente de validación humana', '',
        f'**Eventos procesados:** {len(events)}',
        f'**Equipos:** {", ".join(hosts) if hosts else "No determinados"}',
        f'**Hallazgos:** {len(findings)}', '',
        '## Alcance y limitaciones', '',
        'Esta herramienta analiza reglas y correlaciones sobre exportaciones de eventos y snapshots. '
        'La ausencia de hallazgos NO demuestra ausencia de compromiso. '
        'Los snapshots reflejan el estado al momento de recolección, no un historial. '
        'La correlación temporal NO prueba causalidad.', '',
        '## Hallazgos', '',
    ]
    for f in findings:
        lines += [f"### {f['id']} — {f['title']}", '',
                  f"- **Severidad orientativa:** {f['severity']}",
                  f"- **Categoría:** {f['category']}",
                  f"- **Estado:** {f['status']}",
                  f"- **Descripción:** {f['description']}",
                  f"- **Recomendación:** {f['recommendation']}",
                  f"- **Coincidencias:** {f['evidence_count']}",
                  '- **Referencias a evidencia:**']
        for ev in f['evidence'][:8]:
            lines.append(f"  - `{_safe(ev.get('source', ''))}` — {ev.get('timestamp','')} / Host {ev.get('computer','')} / "
                         f"Evento {ev.get('event_id','')} / Registro {ev.get('record_id','')}")
        lines.append('')
    if ai:
        lines += ['## Hipótesis auxiliares generadas por IA (NO VERIFICADAS)', '',
                  ai['disclaimer'], '', ai['resumen'], '', '**Hipótesis:**', '']
        lines += ['- ' + s for s in ai['hipotesis']]
        lines += ['', '**Validaciones:**', ''] + ['- ' + s for s in ai['validaciones']]
        lines.append('')
    lines += ['## Manifiesto SHA-256 de entradas', '']
    for src in manifest:
        lines += [f"- `{_safe(src['path'])}` — `{src['sha256']}`"]
    lines += ['', '## Advertencias y cobertura', '']
    lines += ['- ' + w for w in warnings] if warnings else ['- Sin advertencias del procesador.']
    lines += ['', '## Próximos pasos', '',
              'Validar manualmente los hallazgos prioritarios, preservar la evidencia original, '
              'correlacionar con procesos, usuarios y fuentes no ingeridas, y documentar las conclusiones.']
    (dst / 'informe.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    # Portable HTML without dependencies: no inline scripts, no external resources.
    out = ['<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">',
           '<title>Informe DFIR preliminar</title>',
           '<style>body{font:15px/1.5 system-ui,sans-serif;color:#1d2939;max-width:1050px;margin:32px auto;padding:0 20px}h1{border-bottom:4px solid #ed7d31;padding-bottom:16px}h2{color:#145a82}section{border:1px solid #ddd;border-radius:8px;margin:15px 0;padding:16px}small{color:#556}.tag{font-weight:700;background:#fff1e6;padding:4px 8px;border-radius:5px}code{overflow-wrap:anywhere}li{margin:4px 0}</style>',
           '</head><body><h1>Informe forense automatizado — PRELIMINAR</h1>',
           f'<p><strong>{len(events)} eventos</strong> / <strong>{len(findings)} hallazgos</strong> / <strong>{len(hosts)} equipos</strong></p>',
           '<p>Validación humana obligatoria. No se afirma compromiso confirmado.</p><h2>Hallazgos</h2>']
    for f in findings:
        out.append('<section><h3>' + html.escape(f['id'] + ' — ' + f['title']) + '</h3>')
        out.append('<p class="tag">' + html.escape(f['severity'].upper()) + '</p>')
        out.append('<p>' + html.escape(f['description']) + '</p><p><strong>Validación:</strong> ' + html.escape(f['recommendation']) + '</p>')
        out.append('<ul>')
        for ev in f['evidence'][:8]:
            out.append('<li><small>' + html.escape(f"{ev.get('timestamp','')} — {ev.get('computer','')} — EventID={ev.get('event_id','')} RecordID={ev.get('record_id','')} — {ev.get('source','')}") + '</small></li>')
        out.append('</ul></section>')
    if ai:
        out.append('<h2>IA local — hipótesis no verificadas</h2><p>' + html.escape(ai['resumen']) + '</p>')
        for h in ai['hipotesis']:
            out.append('<p>• ' + html.escape(h) + '</p>')
    if warnings:
        out.append('<h2>Advertencias</h2><ul>')
        for w in warnings:
            out.append('<li>' + html.escape(w) + '</li>')
        out.append('</ul>')
    out.append('</body></html>')
    (dst / 'informe.html').write_text('\n'.join(out), encoding='utf-8')
    return data


def _make_docx(path, data):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.8)
    sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.85)
    sec.right_margin = Inches(.85)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(9)
    styles['Heading 1'].font.color.rgb = RGBColor(16, 78, 116)
    styles['Heading 2'].font.color.rgb = RGBColor(16, 78, 116)
    title = doc.add_heading('ANÁLISIS FORENSE AUTOMATIZADO', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('REPORTE PRELIMINAR · SUJETO A VALIDACIÓN HUMANA')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor(213, 102, 36)
    doc.add_paragraph('Fecha UTC: ' + _safe(data['generated_utc']))
    doc.add_paragraph(f"Eventos procesados: {data['event_count']} · Equipos: {len(data['hosts'])} · Hallazgos: {len(data['findings'])}")
    doc.add_heading('1. Alcance y limitaciones', level=1)
    doc.add_paragraph('Se aplicaron detecciones deterministas sobre registros procesados y snapshots de configuración. '
                      'No se ha demostrado compromiso por el mero resultado de una regla. La ausencia de alertas '
                      'no es prueba de ausencia de incidentes. Los snapshots son mediciones puntuales.')
    doc.add_heading('2. Resumen de hallazgos', level=1)
    if data['findings']:
        t = doc.add_table(rows=1, cols=3)
        t.style = 'Light Shading Accent 1'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(t.rows[0].cells, ['ID', 'Severidad', 'Título']):
            cell.text = text
        for f in data['findings']:
            cells = t.add_row().cells
            for cell, text in zip(cells, [f['id'], f['severity'].upper(), f['title'][:110]]):
                cell.text = _safe(text)
    else:
        doc.add_paragraph('Sin hallazgos según las reglas aplicadas; la cobertura de fuentes debe comprobarse.')
    doc.add_heading('3. Hallazgos técnicos', level=1)
    for f in data['findings']:
        doc.add_heading(f"{f['id']} | {f['title']}", level=2)
        doc.add_paragraph('Severidad: ' + f['severity'].upper() + '   |   Estado: Pendiente de validación')
        doc.add_paragraph('Observación: ' + _safe(f['description']))
        doc.add_paragraph('Validación recomendada: ' + _safe(f['recommendation']))
        doc.add_paragraph('Referencias (' + str(f['evidence_count']) + ' coincidencias; hasta 8 ejemplos):')
        for ev in f['evidence'][:8]:
            doc.add_paragraph(_safe(f"{ev.get('timestamp','')} · {ev.get('computer','')} · EventID={ev.get('event_id','')} · "
                                     f"RecordID={ev.get('record_id','')} · {ev.get('source','')}")[:440], style='List Bullet')
    next_section = 4
    if data['ai_review']:
        a = data['ai_review']
        doc.add_heading('4. Análisis auxiliar de IA — NO VERIFICADO', level=1)
        doc.add_paragraph(a['disclaimer'])
        doc.add_paragraph(_safe(a['resumen']))
        for h in a['hipotesis']:
            doc.add_paragraph(_safe(h), style='List Bullet')
        doc.add_heading('Validaciones propuestas', level=2)
        for v in a['validaciones']:
            doc.add_paragraph(_safe(v), style='List Bullet')
        next_section = 5
    doc.add_heading(f'{next_section}. Integridad de las entradas', level=1)
    for item in data['source_manifest']:
        doc.add_paragraph(_safe(item['path'])[:120] + ' | SHA-256: ' + item['sha256'], style='List Bullet')
    doc.add_heading(f'{next_section + 1}. Advertencias y cobertura', level=1)
    for w in data['warnings'] or ['Sin advertencias declaradas por el procesador.']:
        doc.add_paragraph(_safe(w)[:500], style='List Bullet')
    footer = sec.footer.paragraphs[0]
    footer.text = 'DFIR AI · Entradas identificadas por SHA-256 · Informe preliminar'
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)
